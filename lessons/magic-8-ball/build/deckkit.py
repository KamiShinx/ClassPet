# -*- coding: utf-8 -*-
"""Builder utilities: append slides to the existing Hebrew lesson deck,
reusing its exact typography, palette and page furniture."""
import copy, os
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Emu, Inches

# ---- the deck's own design tokens -----------------------------------------
INK      = "2C3687"     # body navy
INK2     = "2C36A6"     # brighter navy
PINK     = "E91D63"     # titles / emphasis
TEAL     = "17B3A9"
GREY     = "374151"
GOLD     = "B18E35"     # App Inventor event blocks
GOLD_L   = "E5BE58"
PURPLE   = "704778"     # App Inventor method blocks
GREEN    = "32DE84"
YEL      = "FFE599"
F_BODY   = "Alef"
F_TITLE  = "Open Sans ExtraBold"
F_CODE   = "Source Code Pro"

# ---- geometry copied from the existing slides (EMU) ------------------------
TITLE_FULL = (1325880, 372500, 7506420, 1096200)  # clears the master dashes
BODY_FULL  = (311700, 1468825, 8520600, 2967900)
TITLE_HALF = (311700, 372500, 3996000, 1096200)
BODY_HALF  = (64000, 1468825, 4243000, 3100000)
IMG_HALF   = (4462780, 530000, 4535000, 4280000)
SLIDE_W, SLIDE_H = 9144000, 5143500

TUT = "tutorial"          # drop the real tutorial screenshots here
FALLBACK = "img"


def art(name, fallback=None):
    """Prefer the original tutorial media; fall back to a generated stand-in."""
    for d in (TUT, os.path.join(TUT, "crop")):
        for ext in ("", ".png", ".jpg", ".gif"):
            p = os.path.join(d, name + ext)
            if os.path.exists(p):
                return p, True
    if fallback:
        p = os.path.join(FALLBACK, fallback + ".png")
        if os.path.exists(p):
            return p, False
    return None, False


# ---------------------------------------------------------------- runs ----
def R(t, color=INK, sz=1400, b=False, i=False, font=F_BODY, hl=None, u=False):
    return dict(t=t, color=color, sz=sz, b=b, i=i, font=font, hl=hl, u=u)


def code(t, color=PINK, sz=1300, b=True):
    return R(t, color=color, sz=sz, b=b, font=F_CODE)


def _run_xml(r):
    a = 'lang="iw" sz="%d"' % r["sz"]
    if r["b"]: a += ' b="1"'
    if r["i"]: a += ' i="1"'
    if r["u"]: a += ' u="sng"'
    inner = '<a:solidFill><a:srgbClr val="%s"/></a:solidFill>' % r["color"]
    if r["hl"]:
        inner += '<a:highlight><a:srgbClr val="%s"/></a:highlight>' % r["hl"]
    f = r["font"]
    inner += ('<a:latin typeface="%s"/><a:ea typeface="%s"/>'
              '<a:cs typeface="%s"/><a:sym typeface="%s"/>' % (f, f, f, f))
    txt = (r["t"].replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))
    sp = ' xml:space="preserve"' if txt != txt.strip() else ""
    return '<a:r><a:rPr %s>%s</a:rPr><a:t%s>%s</a:t></a:r>' % (a, inner, sp, txt)


def _ppr(kind="p", before=0, after=0, line=90, align="r", rtl=True,
         bullet_color=INK, sz=1400):
    rt = ' rtl="1"' if rtl else ' rtl="0"'
    if kind == "b":
        head = ('<a:pPr indent="-317500" lvl="0" marL="457200"%s algn="%s">' % (rt, align))
        bul = ('<a:buClr><a:srgbClr val="%s"/></a:buClr><a:buSzPts val="%d"/>'
               '<a:buFont typeface="%s"/><a:buChar char="-"/>' % (bullet_color, sz, F_BODY))
    else:
        head = '<a:pPr indent="0" lvl="0" marL="0"%s algn="%s">' % (rt, align)
        bul = '<a:buNone/>'
    return (head + '<a:lnSpc><a:spcPct val="%d000"/></a:lnSpc>'
            '<a:spcBef><a:spcPts val="%d"/></a:spcBef>'
            '<a:spcAft><a:spcPts val="%d"/></a:spcAft>' % (line, before, after)
            + bul + '</a:pPr>')


def para(runs, kind="p", before=0, after=0, line=90, align="r", rtl=True,
         bullet_color=INK):
    if isinstance(runs, dict):
        runs = [runs]
    sz = runs[0]["sz"] if runs else 1400
    xml = ('<a:p %s>' % nsdecls('a')) + _ppr(kind, before, after, line, align, rtl,
                                             bullet_color, sz)
    xml += "".join(_run_xml(r) for r in runs)
    xml += '</a:p>'
    return parse_xml(xml)


# --------------------------------------------------------------- slides ----
class Deck:
    def __init__(self, prs, template_idx=15):
        self.prs = prs
        self.layout = prs.slide_masters[0].slide_layouts[2]   # TITLE_AND_BODY
        self.tpl = prs.slides[template_idx]                   # furniture donor
        self._furniture = self._collect(self.tpl)

    def _collect(self, s):
        """Background picture + the four corner marks, in z-order."""
        out = []
        for sh in s.shapes:
            if sh.shape_type == 13:      # PICTURE
                out.append(sh._element)
        return out

    def _clone(self, slide, el, index=None):
        new = copy.deepcopy(el)
        for node in new.iter():
            for att in (qn('r:embed'), qn('r:link')):
                if att in node.attrib:
                    rel = self.tpl.part.rels[node.attrib[att]]
                    node.attrib[att] = slide.part.relate_to(rel.target_part, rel.reltype)
        tree = slide.shapes._spTree
        if index is None:
            tree.append(new)
        else:
            tree.insert(index, new)
        return new

    def new(self):
        s = self.prs.slides.add_slide(self.layout)
        for ph in list(s.placeholders):        # start from a clean canvas
            ph._element.getparent().remove(ph._element)
        self._clone(s, self._furniture[0], index=2)      # background, behind all
        for el in self._furniture[1:]:                   # corner marks, on top
            self._clone(s, el)
        return s

    # -- text -------------------------------------------------------------
    def box(self, slide, geo, paras, anchor="t", autofit=False, wrap=True):
        l, t, w, h = geo
        tb = slide.shapes.add_textbox(Emu(l), Emu(t), Emu(w), Emu(h))
        tf = tb.text_frame
        bp = tf._txBody.find(qn('a:bodyPr'))
        bp.set('anchor', anchor); bp.set('anchorCtr', '0')
        bp.set('lIns', '91425'); bp.set('rIns', '91425')
        bp.set('tIns', '45700'); bp.set('bIns', '45700')
        bp.set('wrap', 'square' if wrap else 'none')
        bp.set('spcFirstLastPara', '1')
        for ch in list(bp):
            bp.remove(ch)
        bp.append(parse_xml('<a:%s %s/>' % ('normAutofit' if autofit else 'noAutofit',
                                            nsdecls('a'))))
        for p in list(tf._txBody.findall(qn('a:p'))):
            tf._txBody.remove(p)
        for p in paras:
            tf._txBody.append(p)
        if not tf._txBody.findall(qn('a:p')):
            tf._txBody.append(parse_xml('<a:p %s/>' % nsdecls('a')))
        return tb

    def title(self, slide, text, half=False, sz=None, color=PINK):
        geo = TITLE_HALF if half else TITLE_FULL
        if sz is None:                       # keep clear of the master's dashes
            n = len(text)
            sz = 2800 if n <= 24 else 2600 if n <= 34 else 2400
        return self.box(slide, geo,
                        [para(R(text, color=color, sz=sz, font=F_TITLE), line=95)],
                        anchor="b", autofit=True)

    # -- pictures ---------------------------------------------------------
    def pic(self, slide, path, box, mode="fit", border=None, max_dpi=None):
        """Place `path` inside box=(l,t,w,h) EMU, preserving aspect ratio.
        max_dpi caps the enlargement so equivalent art keeps one scale."""
        from PIL import Image
        l, t, w, h = box
        iw, ih = Image.open(path).size
        k = min(w / iw, h / ih) if mode == "fit" else max(w / iw, h / ih)
        if max_dpi:
            k = min(k, 914400.0 / max_dpi)
        nw, nh = int(iw * k), int(ih * k)
        nl, nt = int(l + (w - nw) / 2), int(t + (h - nh) / 2)
        p = slide.shapes.add_picture(path, Emu(nl), Emu(nt), Emu(nw), Emu(nh))
        if border:
            p.line.color.rgb = border[0]; p.line.width = border[1]
        return p

    def rrect(self, slide, box, fill=None, line=None, lw=12700, radius=0.08):
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.dml.color import RGBColor
        l, t, w, h = box
        sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                    Emu(l), Emu(t), Emu(w), Emu(h))
        sh.adjustments[0] = radius
        if fill:
            sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(fill)
        else:
            sh.fill.background()
        if line:
            sh.line.color.rgb = RGBColor.from_string(line); sh.line.width = lw
        else:
            sh.line.fill.background()
        sh.shadow.inherit = False
        # keep exactly one empty <a:p>: PowerPoint rejects an empty <p:txBody>
        ps = sh.text_frame._txBody.findall(qn('a:p'))
        for extra in ps[1:]:
            sh.text_frame._txBody.remove(extra)
        return sh
